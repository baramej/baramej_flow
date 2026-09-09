# -*- coding: utf-8 -*-
"""Reads an uploaded .xlsx file (an Odoo ir.attachment) into a list of row dicts,
keyed by the header row — e.g. feeding a knowledge base / reference spreadsheet
into context for an AI node's prompt, or bulk-processing rows into Create Record
actions via a future Phase 4 Loop node.
"""
import base64
import io
import logging

from odoo import _
from odoo.exceptions import UserError

from ..node_registry import register_node
from ..utils import render_template
from .base import BaseNodeExecutor

_logger = logging.getLogger(__name__)
_MAX_ROWS = 5000  # safety cap so a huge spreadsheet doesn't blow up the execution context


@register_node('read_excel')
class ReadExcelExecutor(BaseNodeExecutor):
    """Expected config_json:
        {
            "attachment_id": 42,
            "sheet_name": "Sheet1",
            "has_header_row": true,
            "output_key": "excel_rows"
        }

    "attachment_id" can also be a {{dotted.path}} template resolving to an id
    (e.g. from an earlier node that just uploaded a file) — it's rendered through
    the same templating helper as every other node before being used.

    With has_header_row=true (default), the first row becomes each row dict's
    keys (e.g. {"Ticket ID": 101, "Priority": "High"}). With false, rows are plain
    lists of cell values instead.
    """

    def execute(self):
        config = self.node.get_config()
        attachment_id_raw = config.get('attachment_id')
        if not attachment_id_raw:
            raise UserError(_('Read Excel node "%s" is missing "attachment_id" in its configuration.') %
                             self.node.name)

        attachment_id = render_template(str(attachment_id_raw), self.context)
        try:
            attachment_id = int(attachment_id)
        except (TypeError, ValueError):
            raise UserError(_('Read Excel node "%s" resolved "attachment_id" to "%s", which is not a valid '
                               'id.') % (self.node.name, attachment_id))

        attachment = self.env['ir.attachment'].sudo().browse(attachment_id)
        if not attachment.exists():
            raise UserError(_('Read Excel node "%s" references attachment id %s, which does not exist.') % (
                self.node.name, attachment_id))

        try:
            import openpyxl
        except ImportError:
            raise UserError(_('openpyxl is not installed in this Odoo environment. Run: pip install openpyxl'))

        try:
            file_bytes = base64.b64decode(attachment.datas or b'')
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001 - surface openpyxl's own message verbatim
            raise UserError(_('Read Excel node "%s" could not open "%s" as an Excel file: %s') % (
                self.node.name, attachment.name, exc))

        sheet_name = config.get('sheet_name')
        try:
            sheet = workbook[sheet_name] if sheet_name else workbook.active
        except KeyError:
            raise UserError(_('Read Excel node "%s": sheet "%s" not found in "%s". Available sheets: %s') % (
                self.node.name, sheet_name, attachment.name, ', '.join(workbook.sheetnames)))

        has_header_row = config.get('has_header_row', True)
        rows = []
        header = None

        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0 and has_header_row:
                header = [str(c) if c is not None else ('col_%d' % idx) for idx, c in enumerate(row)]
                continue
            if len(rows) >= _MAX_ROWS:
                _logger.warning('Read Excel node %s: truncated at %s rows (file has more).',
                                 self.node.name, _MAX_ROWS)
                break
            if header:
                rows.append(dict(zip(header, row)))
            else:
                rows.append(list(row))

        output_key = config.get('output_key', 'excel_rows')
        new_context = dict(self.context)
        new_context[output_key] = rows
        return {'output': new_context}
